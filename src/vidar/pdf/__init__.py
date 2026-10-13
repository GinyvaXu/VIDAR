"""PDF 渲染入口：DocData 数据装载 + 模板分发。"""

from __future__ import annotations

from pathlib import Path

from ..errors import DependencyError, VidarError
from ..frames import ExtractedFrame
from ..models import (
    AsrResult,
    ParagraphResult,
    RefineResult,
    SummaryResult,
    TranslateResult,
    VideoMeta,
)

# 字号档位 → 缩放系数（compact 小字一页能放更多内容）
SCALE_PRESETS = {"compact": 0.85, "standard": 1.0, "large": 1.18}
POINT_FIGURE_MODES = ("all", "sparse", "off")


def resolve_scale(value: object, default: float = 1.0) -> float:
    """档位名（compact/standard/large）或数字字符串 → 缩放系数（0.6–1.5）。"""
    if isinstance(value, bool):
        return default
    if isinstance(value, (int, float)):
        factor = float(value)
    else:
        text = str(value or "").strip().lower()
        if not text:
            return default
        if text in SCALE_PRESETS:
            factor = SCALE_PRESETS[text]
        else:
            try:
                factor = float(text)
            except ValueError:
                return default
    return min(1.5, max(0.6, factor))


class DocData:
    """PDF / DOCX 渲染所需的全部数据（从 work 目录产物组装）。"""

    output_path: Path  # 由渲染函数注入

    def __init__(
        self,
        *,
        meta: VideoMeta,
        summary: SummaryResult,
        refine: RefineResult,
        asr: AsrResult,
        frames: list[ExtractedFrame],
        images_dir: Path,
        mode: str = "full",
        style: str = "newspaper",
        scale: str | float = "standard",
        point_figures: str = "all",
        point_thumb_mm: float = 28.0,
        translation: TranslateResult | None = None,
        bilingual: bool = True,
        paragraphs: ParagraphResult | None = None,
        ads_note: str = "",
    ) -> None:
        self.meta = meta
        self.summary = summary
        self.refine = refine
        self.asr = asr
        self.frames = frames
        self.images_dir = images_dir
        self.mode = mode if mode in {"full", "brief"} else "full"
        self.style = style if style in {"newspaper", "paper"} else "newspaper"
        self.scale = resolve_scale(scale, 1.0)
        self.point_figures = point_figures if point_figures in POINT_FIGURE_MODES else "all"
        self.point_thumb_mm = float(point_thumb_mm) if point_thumb_mm else 28.0
        self.translation = translation
        self.bilingual = bool(bilingual) and translation is not None
        self.paragraphs = paragraphs
        self.ads_note = ads_note
        # 译文映射始终构建：关闭双语时正文仍用译文（文档语言 = 目标语言）
        self._translations: dict[int, str] = {}
        if translation is not None:
            self._translations = {
                seg.idx: seg.translated_text
                for seg in translation.segments
                if seg.translated_text.strip()
            }

    # ------------------------------------------------------------------ #
    def chapter_paragraphs(self, chapter) -> list[tuple[float, str, str]]:
        """章内段落（按时间归属）：(时间, 正文, 原文)。

        优先使用语段编辑（paragraphs.json）的结果：一个自然段一行；
        旧任务（无语段产物）回退为逐句渲染。有译文时正文=译文，原文仅在双语排版时附带。
        """
        start = chapter.start
        end = chapter.end if chapter.end is not None else float("inf")

        if self.paragraphs is not None and self.paragraphs.paragraphs:
            rows: list[tuple[float, str, str]] = []
            for para in self.paragraphs.paragraphs:
                if not (start <= para.start < end):
                    continue
                source = para.source_text if self.bilingual else ""
                rows.append((para.start, para.text, source))
            # 语段产物存在时不做逐句回退：跨章段落已在起始章展示，
            # 本章若无段起点宁可留白，也不重复渲染（避免同一段文字跨章出现两次）
            return rows

        rows = []
        for segment in self.refine.segments:
            if not (start <= segment.start < end):
                continue
            translated = self._translations.get(segment.idx, "")
            if translated:
                source = segment.refined_text if self.bilingual else ""
                rows.append((segment.start, translated, source))
            else:
                rows.append((segment.start, segment.refined_text, ""))
        return rows

    def paragraph_source(self, idx: int) -> str:
        """指定段落的译文（无则空串）——供逐字稿/SRT 双语渲染。"""
        return self._translations.get(idx, "")

    def chapter_figures(self, chapter) -> list[tuple[float, str, Path]]:
        """章内配图（时间, 图注, 文件路径）。"""
        result: list[tuple[float, str, Path]] = []
        for hint in chapter.figures:
            for frame in self.frames:
                if abs(frame.start - hint.start) < 0.5:
                    path = self.images_dir / frame.filename
                    if path.exists():
                        result.append((hint.start, frame.caption or hint.caption, path))
                    break
        return result

    def chapter_quote(self, chapter) -> str:
        """章节金句（与小结重复时自动去重，避免排版冗余）。"""
        quote = (chapter.quote or "").strip()
        summary = (chapter.summary or "").strip()
        if quote and summary and (quote in summary or summary in quote):
            return ""
        return quote

    def point_thumb(self, order: int) -> Path | None:
        """第 order 个要点（1 起）的截图；按「密度」档位筛选（off / sparse 隔一取一）。"""
        if self.point_figures == "off":
            return None
        if self.point_figures == "sparse" and order % 2 == 0:
            return None
        for frame in self.frames:
            if frame.kind == "point" and frame.point_index == order:
                path = self.images_dir / frame.filename
                if path.exists():
                    return path
        return None

    def cover_image(self) -> Path | None:
        cover = self.images_dir / "cover.jpg"
        return cover if cover.exists() else None

    @property
    def keywords(self) -> list[str]:
        return [tag for tag in self.meta.tags if tag][:6]


def render_pdf(data: DocData, output_path: Path) -> Path:
    """按 style 渲染 PDF，返回产物路径。"""
    try:
        import reportlab  # noqa: F401
    except ImportError as exc:  # pragma: no cover
        raise DependencyError("缺少 PDF 依赖 reportlab", hint="执行 uv sync 安装") from exc

    data.output_path = output_path
    if data.style == "paper":
        from .paper import build_paper

        return build_paper(data)

    from .newspaper import build_newspaper

    return build_newspaper(data)


def render_docx(data: DocData, output_path: Path) -> Path:
    """渲染 Word 文档（.docx），返回产物路径。"""
    try:
        import docx  # noqa: F401
    except ImportError as exc:  # pragma: no cover
        raise DependencyError("缺少 Word 依赖 python-docx", hint="执行 uv sync 安装") from exc

    from .docx_export import build_docx

    return build_docx(data, output_path)


__all__ = ["DocData", "render_pdf", "render_docx", "resolve_scale", "SCALE_PRESETS", "VidarError"]
