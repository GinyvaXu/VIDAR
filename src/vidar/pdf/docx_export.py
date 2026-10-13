"""Word 导出（.docx）：标题 / 摘要 / 要点 / 章节（金句 + 配图 + 全文）。"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .. import __version__
from ..utils.text import format_clock
from . import DocData

EAST_BODY = "宋体"
EAST_HEAD = "微软雅黑"
LATIN_BODY = "Times New Roman"
LATIN_HEAD = "Microsoft YaHei"
GRAY = RGBColor(0x5F, 0x5F, 0x5F)
IMAGE_WIDTH_CM = 13.0
IMAGE_MAX_HEIGHT_CM = 21.5  # A4 文本区高度上限（含图注留白）


def _add_picture_fitted(document, path, *, width_cm: float) -> None:  # noqa: ANN001
    """插图：固定宽度；竖版截图（高宽比过大）按文本区高度封顶，避免溢出页面。"""
    from PIL import Image

    with Image.open(path) as image:
        ratio = image.height / max(image.width, 1)
    width = Cm(width_cm)
    height = Cm(width_cm * ratio)
    if width_cm * ratio > IMAGE_MAX_HEIGHT_CM:
        height = Cm(IMAGE_MAX_HEIGHT_CM)
        width = Cm(IMAGE_MAX_HEIGHT_CM / ratio)
    document.add_picture(str(path), width=width, height=height)
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _set_fonts(element, *, east: str, latin: str) -> None:  # noqa: ANN001
    rpr = element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)


def _run(
    paragraph,
    text: str,
    *,
    east: str = EAST_BODY,
    latin: str = LATIN_BODY,
    size: float = 10.5,
    bold: bool = False,
    color: RGBColor | None = None,
):  # noqa: ANN001, ANN202
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = latin
    if color is not None:
        run.font.color.rgb = color
    _set_fonts(run._element, east=east, latin=latin)
    return run


def _heading(document, text: str, *, size: float = 13.5) -> None:  # noqa: ANN001
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    _run(paragraph, text, east=EAST_HEAD, latin=LATIN_HEAD, size=size, bold=True)


def _body(
    document,
    text: str,
    *,
    size: float = 10.5,
    gray: bool = False,
    indent: float = 0.0,
    first_line: float = 0.0,
):  # noqa: ANN001, ANN202
    paragraph = document.add_paragraph()
    if indent:
        paragraph.paragraph_format.left_indent = Cm(indent)
    if first_line:
        paragraph.paragraph_format.first_line_indent = Pt(first_line)
    _run(paragraph, text, size=size, color=GRAY if gray else None)
    return paragraph


def build_docx(data: DocData, output_path: Path) -> Path:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = LATIN_BODY
    _set_fonts(normal.element, east=EAST_BODY, latin=LATIN_BODY)

    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)

    meta = data.meta
    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(
        title_paragraph,
        meta.title,
        east=EAST_HEAD,
        latin=LATIN_HEAD,
        size=18,
        bold=True,
    )

    info_paragraph = document.add_paragraph()
    info_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    published = meta.pubdate.strftime("%Y-%m-%d") if meta.pubdate else "—"
    _run(
        info_paragraph,
        f"作者：{meta.up_name or '—'}　整理：VIDAR v{__version__}",
        size=9,
        color=GRAY,
    )
    meta_paragraph = document.add_paragraph()
    meta_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(
        meta_paragraph,
        f"{published} · 时长 {format_clock(meta.duration_sec or 0)} · {meta.bvid} · {meta.url}",
        size=9,
        color=GRAY,
    )

    _heading(document, "摘要")
    if data.summary.one_liner:
        _body(document, data.summary.one_liner)
    for index, point in enumerate(data.summary.key_points[:5], start=1):
        _body(document, f"{index}. {point.text}")
    if data.keywords:
        _body(document, f"关键词：{' / '.join(data.keywords)}", size=9, gray=True)

    figure_number = 1
    for chapter in data.summary.chapters:
        _heading(
            document,
            f"{chapter.index + 1}　{chapter.title}　[{format_clock(chapter.start)}]",
        )
        if chapter.summary:
            _body(document, chapter.summary, size=10, gray=True)
        quote = data.chapter_quote(chapter)
        if quote:
            _body(document, f"— {quote}", size=10, gray=True, indent=0.6)

        for start, caption, path in data.chapter_figures(chapter):
            _add_picture_fitted(document, path, width_cm=IMAGE_WIDTH_CM)
            caption_paragraph = document.add_paragraph()
            caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            label = (
                f"图 {figure_number}：{caption}（{format_clock(start)}）"
                if caption
                else f"图 {figure_number}（{format_clock(start)}）"
            )
            _run(caption_paragraph, label, size=9, color=GRAY)
            figure_number += 1

        if data.mode == "full":
            for seg_start, text, source in data.chapter_paragraphs(chapter):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.first_line_indent = Pt(21)  # 首行缩进两格
                _run(paragraph, text)
                _run(paragraph, f"　[{format_clock(seg_start)}]", size=8, color=GRAY)
                if source:
                    _body(document, source, size=9, gray=True, indent=0.5, first_line=18)

    trailer = document.add_paragraph()
    _run(
        trailer,
        f"本文档由 VIDAR v{__version__} 生成于 {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        size=9,
        color=GRAY,
    )
    if data.ads_note:
        ads_paragraph = document.add_paragraph()
        _run(ads_paragraph, data.ads_note, size=9, color=GRAY)

    document.save(str(output_path))
    return output_path
